from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
import sqlite3
import statistics
import uuid
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo


APP_TZ = ZoneInfo(os.getenv("POS_TIMEZONE", "Asia/Karachi"))
DEFAULT_DB_PATH = Path(os.getenv("POS_DB_PATH", "backend/pos_analytics.sqlite3"))
SCHEMA_VERSION = 1

TRANSACTION_ALIASES = {
    "line_item_id": {"line_item_id", "line id", "line_id", "item_line_id"},
    "transaction_id": {"transaction_id", "transaction id", "txn_id", "receipt_id", "bill_no", "invoice_no"},
    "timestamp": {"timestamp", "time", "datetime", "date_time", "transaction_time", "sold_at"},
    "business_date": {"business_date", "date", "sale_date"},
    "till_id": {"till_id", "till", "counter_id", "register_id"},
    "sku_id": {"sku_id", "sku", "barcode", "item_code", "product_id"},
    "sku_name": {"sku_name", "product", "product_name", "item_name", "description"},
    "category": {"category", "department", "class"},
    "qty": {"qty", "quantity", "units", "items"},
    "unit_price": {"unit_price", "price", "rate", "selling_price"},
    "total_price": {"total_price", "line_total", "amount", "net_amount", "gross_amount"},
    "unit_cost": {"unit_cost", "cost", "cogs", "buying_price"},
    "transaction_type": {"transaction_type", "type", "sale_type", "status"},
}

INVENTORY_ALIASES = {
    "snapshot_at": {"snapshot_at", "timestamp", "stock_at", "as_of", "date"},
    "sku_id": {"sku_id", "sku", "barcode", "item_code", "product_id"},
    "on_hand": {"on_hand", "stock", "stock_on_hand", "qty_on_hand", "available_stock"},
    "on_order": {"on_order", "incoming", "incoming_qty", "ordered_qty"},
    "receipt_due_at": {"receipt_due_at", "due_at", "expected_at", "receipt_date", "arrival_date"},
    "lead_time_days": {"lead_time_days", "lead_time", "lead_days"},
    "safety_stock_units": {"safety_stock_units", "safety_stock", "minimum_stock", "min_stock"},
    "unit_cost": {"unit_cost", "cost", "cogs", "buying_price"},
    "comparison_group": {"comparison_group", "peer_group", "comp_group"},
}

FOOTFALL_ALIASES = {
    "event_id": {"event_id", "id", "crossing_id"},
    "timestamp": {"timestamp", "time", "datetime", "event_time"},
    "direction": {"direction", "dir", "event_type", "in_out"},
    "camera_id": {"camera_id", "camera", "sensor_id", "source_camera"},
    "sequence": {"sequence", "seq", "source_sequence", "frame_sequence"},
}

CALENDAR_ALIASES = {
    "date": {"date", "business_date", "open_date"},
    "is_open": {"is_open", "open", "store_open"},
}

REQUIRED = {
    "transactions": {"transaction_id", "timestamp", "till_id", "sku_id", "sku_name", "category", "qty", "unit_price"},
    "inventory": {"snapshot_at", "sku_id", "on_hand"},
    "footfall": {"timestamp", "direction", "camera_id"},
    "calendar": {"date", "is_open"},
}


def connect(db_path: Path | str = DEFAULT_DB_PATH) -> sqlite3.Connection:
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    initialize(conn)
    return conn


def initialize(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        INSERT OR IGNORE INTO meta(key, value) VALUES ('schema_version', '1');
        CREATE TABLE IF NOT EXISTS imports (
          id TEXT PRIMARY KEY,
          source TEXT NOT NULL,
          kind TEXT NOT NULL,
          coverage_from TEXT,
          coverage_to TEXT,
          complete INTEGER NOT NULL DEFAULT 0,
          row_count INTEGER NOT NULL DEFAULT 0,
          committed_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS transactions (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          source TEXT NOT NULL,
          line_item_id TEXT,
          transaction_id TEXT NOT NULL,
          business_date TEXT NOT NULL,
          timestamp TEXT NOT NULL,
          till_id TEXT NOT NULL,
          sku_id TEXT NOT NULL,
          sku_name TEXT NOT NULL,
          category TEXT NOT NULL,
          qty REAL NOT NULL,
          unit_price REAL NOT NULL,
          total_price REAL NOT NULL,
          unit_cost REAL,
          transaction_type TEXT NOT NULL,
          fingerprint TEXT NOT NULL UNIQUE,
          FOREIGN KEY(import_id) REFERENCES imports(id)
        );
        CREATE INDEX IF NOT EXISTS idx_tx_time ON transactions(timestamp);
        CREATE INDEX IF NOT EXISTS idx_tx_sku ON transactions(sku_id);
        CREATE TABLE IF NOT EXISTS inventory_snapshots (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          snapshot_at TEXT NOT NULL,
          sku_id TEXT NOT NULL,
          on_hand REAL NOT NULL,
          unit_cost REAL,
          comparison_group TEXT,
          lead_time_days INTEGER,
          safety_stock_units REAL,
          fingerprint TEXT NOT NULL UNIQUE,
          FOREIGN KEY(import_id) REFERENCES imports(id)
        );
        CREATE INDEX IF NOT EXISTS idx_inventory_sku_time ON inventory_snapshots(sku_id, snapshot_at);
        CREATE TABLE IF NOT EXISTS inventory_receipts (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          sku_id TEXT NOT NULL,
          quantity REAL NOT NULL,
          receipt_due_at TEXT,
          fingerprint TEXT NOT NULL UNIQUE,
          FOREIGN KEY(import_id) REFERENCES imports(id)
        );
        CREATE TABLE IF NOT EXISTS footfall_events (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          event_id TEXT NOT NULL,
          timestamp TEXT NOT NULL,
          direction TEXT NOT NULL,
          camera_id TEXT NOT NULL,
          fingerprint TEXT NOT NULL UNIQUE,
          FOREIGN KEY(import_id) REFERENCES imports(id)
        );
        CREATE INDEX IF NOT EXISTS idx_footfall_time ON footfall_events(timestamp);
        CREATE TABLE IF NOT EXISTS camera_coverage (
          id TEXT PRIMARY KEY,
          import_id TEXT NOT NULL,
          camera_id TEXT NOT NULL,
          coverage_from TEXT NOT NULL,
          coverage_to TEXT NOT NULL,
          FOREIGN KEY(import_id) REFERENCES imports(id)
        );
        CREATE TABLE IF NOT EXISTS store_calendar (
          date TEXT PRIMARY KEY,
          is_open INTEGER NOT NULL
        );
        CREATE TABLE IF NOT EXISTS analysis_runs (
          id TEXT PRIMARY KEY,
          dataset_version TEXT NOT NULL,
          date_from TEXT NOT NULL,
          date_to TEXT NOT NULL,
          horizon INTEGER NOT NULL,
          status TEXT NOT NULL,
          created_at TEXT NOT NULL,
          completed_at TEXT,
          result_json TEXT,
          settings_json TEXT NOT NULL DEFAULT '{}'
        );
        CREATE INDEX IF NOT EXISTS idx_runs_lookup ON analysis_runs(date_from, date_to, horizon, status);
        CREATE TABLE IF NOT EXISTS evidence (
          id TEXT PRIMARY KEY,
          run_id TEXT NOT NULL,
          scope_type TEXT NOT NULL,
          scope_id TEXT NOT NULL,
          period_from TEXT NOT NULL,
          period_to TEXT NOT NULL,
          text TEXT NOT NULL,
          payload_json TEXT NOT NULL,
          FOREIGN KEY(run_id) REFERENCES analysis_runs(id)
        );
        CREATE VIRTUAL TABLE IF NOT EXISTS evidence_fts USING fts5(id UNINDEXED, text);
        """
    )
    conn.commit()


def normalize_header(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", header.strip().lower()).strip("_")


def alias_map(kind: str, headers: Iterable[str]) -> tuple[dict[str, str], list[str]]:
    alias_source = {
        "transactions": TRANSACTION_ALIASES,
        "inventory": INVENTORY_ALIASES,
        "footfall": FOOTFALL_ALIASES,
        "calendar": CALENDAR_ALIASES,
    }[kind]
    normalized = {normalize_header(h): h for h in headers}
    mapped: dict[str, str] = {}
    errors: list[str] = []
    for canonical, aliases in alias_source.items():
        matches = [normalized[a] for a in aliases if a in normalized]
        if len(matches) > 1:
            errors.append(f"Ambiguous aliases for {canonical}: {', '.join(matches)}")
        elif matches:
            mapped[canonical] = matches[0]
    missing = sorted(REQUIRED[kind] - set(mapped))
    if missing:
        errors.append(f"Missing required columns: {', '.join(missing)}")
    return mapped, errors


def parse_bool(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y", "open"}


def parse_float(value: Any, field: str, required: bool = True) -> float | None:
    if value is None or str(value).strip() == "":
        if required:
            raise ValueError(f"{field} is required")
        return None
    return float(str(value).replace(",", "").strip())


def parse_int(value: Any, field: str, required: bool = False) -> int | None:
    parsed = parse_float(value, field, required)
    if parsed is None:
        return None
    return int(parsed)


def parse_timestamp(value: Any, field: str = "timestamp") -> datetime:
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field} is required")
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip().replace("Z", "+00:00")
        formats = [
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d",
            "%d/%m/%Y %H:%M:%S",
            "%d/%m/%Y %H:%M",
            "%d/%m/%Y",
        ]
        dt = None
        try:
            dt = datetime.fromisoformat(text)
        except ValueError:
            for fmt in formats:
                try:
                    dt = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
        if dt is None:
            raise ValueError(f"{field} must be an ISO-like date or timestamp")
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=APP_TZ)
    return dt.astimezone(APP_TZ)


def iso(dt: datetime) -> str:
    return dt.astimezone(APP_TZ).isoformat()


def business_date(dt: datetime) -> str:
    return dt.astimezone(APP_TZ).date().isoformat()


def stable_hash(*parts: Any) -> str:
    joined = "|".join("" if p is None else str(p) for p in parts)
    return hashlib.sha256(joined.encode("utf-8")).hexdigest()


@dataclass
class ParsedSheet:
    kind: str
    rows: list[dict[str, Any]]
    errors: list[dict[str, Any]]
    warnings: list[str]
    headers: list[str]


def read_csv_bytes(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(text)))


def read_workbook_bytes(data: bytes) -> dict[str, list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # pragma: no cover - depends on optional package presence
        raise RuntimeError("XLSX support requires openpyxl. Install backend/requirements.txt.") from exc
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets: dict[str, list[dict[str, Any]]] = {}
    for name in workbook.sheetnames:
        canonical = normalize_header(name)
        if canonical not in {"transactions", "inventory", "footfall", "calendar"}:
            continue
        worksheet = workbook[name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            sheets[canonical] = []
            continue
        headers = [str(cell or "").strip() for cell in rows[0]]
        records = []
        for row in rows[1:]:
            records.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
        sheets[canonical] = records
    return sheets


def normalize_transaction(row: dict[str, Any], mapped: dict[str, str], source: str) -> dict[str, Any]:
    ts = parse_timestamp(row.get(mapped["timestamp"]))
    tx_type = str(row.get(mapped.get("transaction_type", ""), "sale") or "sale").strip().lower()
    if tx_type not in {"sale", "refund", "void"}:
        raise ValueError("transaction_type must be sale, refund, or void")
    qty = parse_float(row.get(mapped["qty"]), "qty") or 0.0
    unit_price = parse_float(row.get(mapped["unit_price"]), "unit_price") or 0.0
    total = parse_float(row.get(mapped.get("total_price", "")), "total_price", required=False)
    if total is None:
        total = qty * unit_price
    if tx_type == "refund":
        if qty > 0 and total > 0:
            qty, total = -qty, -total
        elif qty < 0 and total < 0:
            pass
        else:
            raise ValueError("refund qty and total_price signs must both be positive or both be negative")
    if tx_type == "sale" and (qty < 0 or total < 0):
        raise ValueError("sale qty and total_price must be non-negative")
    if tx_type == "void":
        qty = 0.0
        total = 0.0
    tx_id = str(row.get(mapped["transaction_id"]) or "").strip()
    till_id = str(row.get(mapped["till_id"]) or "").strip()
    sku_id = str(row.get(mapped["sku_id"]) or "").strip()
    if not tx_id or not till_id or not sku_id:
        raise ValueError("transaction_id, till_id, and sku_id are required")
    bdate = row.get(mapped.get("business_date", ""))
    bdate_str = parse_timestamp(bdate, "business_date").date().isoformat() if bdate else business_date(ts)
    line_id = str(row.get(mapped.get("line_item_id", "")) or "").strip() or None
    fingerprint = stable_hash("tx", source, line_id or "", bdate_str, till_id, tx_id, sku_id, unit_price, qty, total)
    if line_id:
        fingerprint = stable_hash("tx-line", source, line_id)
    return {
        "id": stable_hash("tx-row", fingerprint),
        "source": source,
        "line_item_id": line_id,
        "transaction_id": tx_id,
        "business_date": bdate_str,
        "timestamp": iso(ts),
        "till_id": till_id,
        "sku_id": sku_id,
        "sku_name": str(row.get(mapped["sku_name"]) or sku_id).strip(),
        "category": str(row.get(mapped["category"]) or "Uncategorized").strip(),
        "qty": qty,
        "unit_price": unit_price,
        "total_price": total,
        "unit_cost": parse_float(row.get(mapped.get("unit_cost", "")), "unit_cost", required=False),
        "transaction_type": tx_type,
        "fingerprint": fingerprint,
    }


def normalize_inventory(row: dict[str, Any], mapped: dict[str, str]) -> tuple[dict[str, Any], dict[str, Any] | None]:
    ts = parse_timestamp(row.get(mapped["snapshot_at"]), "snapshot_at")
    sku_id = str(row.get(mapped["sku_id"]) or "").strip()
    if not sku_id:
        raise ValueError("sku_id is required")
    snapshot = {
        "id": stable_hash("inv", sku_id, iso(ts)),
        "snapshot_at": iso(ts),
        "sku_id": sku_id,
        "on_hand": parse_float(row.get(mapped["on_hand"]), "on_hand") or 0.0,
        "unit_cost": parse_float(row.get(mapped.get("unit_cost", "")), "unit_cost", required=False),
        "comparison_group": str(row.get(mapped.get("comparison_group", "")) or "").strip() or None,
        "lead_time_days": parse_int(row.get(mapped.get("lead_time_days", "")), "lead_time_days"),
        "safety_stock_units": parse_float(row.get(mapped.get("safety_stock_units", "")), "safety_stock_units", required=False),
    }
    snapshot["fingerprint"] = stable_hash("inv-snap", sku_id, snapshot["snapshot_at"], snapshot["on_hand"])
    receipt = None
    on_order = parse_float(row.get(mapped.get("on_order", "")), "on_order", required=False)
    if on_order and on_order > 0:
        due_raw = row.get(mapped.get("receipt_due_at", ""))
        due = iso(parse_timestamp(due_raw, "receipt_due_at")) if due_raw else None
        receipt = {
            "id": stable_hash("receipt", sku_id, on_order, due or "undated"),
            "sku_id": sku_id,
            "quantity": on_order,
            "receipt_due_at": due,
            "fingerprint": stable_hash("receipt", sku_id, on_order, due or "undated"),
        }
    return snapshot, receipt


def normalize_footfall(row: dict[str, Any], mapped: dict[str, str], source: str) -> dict[str, Any]:
    ts = parse_timestamp(row.get(mapped["timestamp"]))
    direction = str(row.get(mapped["direction"]) or "").strip().lower()
    direction = {"entry": "in", "enter": "in", "entered": "in", "exit": "out", "leave": "out", "left": "out"}.get(direction, direction)
    if direction not in {"in", "out"}:
        raise ValueError("direction must be in or out")
    camera_id = str(row.get(mapped["camera_id"]) or "").strip()
    if not camera_id:
        raise ValueError("camera_id is required")
    event_id = str(row.get(mapped.get("event_id", "")) or "").strip()
    if not event_id:
        sequence = str(row.get(mapped.get("sequence", "")) or "").strip()
        if not sequence:
            raise ValueError("event_id is required unless a stable source sequence is provided")
        event_id = stable_hash("footfall-seq", source, camera_id, sequence)
    fingerprint = stable_hash("footfall", source, event_id)
    return {
        "id": stable_hash("footfall-row", fingerprint),
        "event_id": event_id,
        "timestamp": iso(ts),
        "direction": direction,
        "camera_id": camera_id,
        "fingerprint": fingerprint,
    }


def normalize_calendar(row: dict[str, Any], mapped: dict[str, str]) -> dict[str, Any]:
    date = parse_timestamp(row.get(mapped["date"]), "date").date().isoformat()
    return {"date": date, "is_open": 1 if parse_bool(row.get(mapped["is_open"])) else 0}


def parse_sheet(kind: str, rows: list[dict[str, Any]], source: str) -> ParsedSheet:
    headers = list(rows[0].keys()) if rows else []
    mapped, header_errors = alias_map(kind, headers)
    parsed: list[dict[str, Any]] = []
    row_errors: list[dict[str, Any]] = []
    warnings: list[str] = []
    if header_errors:
        return ParsedSheet(kind, [], [{"row": 0, "errors": header_errors}], warnings, headers)
    for idx, row in enumerate(rows, start=2):
        try:
            if kind == "transactions":
                parsed.append(normalize_transaction(row, mapped, source))
            elif kind == "inventory":
                snapshot, receipt = normalize_inventory(row, mapped)
                parsed.append({"snapshot": snapshot, "receipt": receipt})
                if receipt and receipt["receipt_due_at"] is None:
                    warnings.append(f"Row {idx}: on_order has no receipt date and will not prevent projected stockouts")
            elif kind == "footfall":
                parsed.append(normalize_footfall(row, mapped, source))
            elif kind == "calendar":
                parsed.append(normalize_calendar(row, mapped))
        except Exception as exc:
            row_errors.append({"row": idx, "errors": [str(exc)]})
    return ParsedSheet(kind, parsed, row_errors, sorted(set(warnings)), headers)


def templates() -> dict[str, Any]:
    return {
        "timezone": str(APP_TZ),
        "currency": "PKR",
        "sheets": {
            "transactions": ["line_item_id", "transaction_id", "timestamp", "business_date", "till_id", "sku_id", "sku_name", "category", "qty", "unit_price", "total_price", "unit_cost", "transaction_type"],
            "inventory": ["snapshot_at", "sku_id", "on_hand", "on_order", "receipt_due_at", "lead_time_days", "safety_stock_units", "unit_cost", "comparison_group"],
            "footfall": ["event_id", "timestamp", "direction", "camera_id", "sequence"],
            "calendar": ["date", "is_open"],
        },
        "aliases": {
            "transactions": {k: sorted(v) for k, v in TRANSACTION_ALIASES.items()},
            "inventory": {k: sorted(v) for k, v in INVENTORY_ALIASES.items()},
            "footfall": {k: sorted(v) for k, v in FOOTFALL_ALIASES.items()},
            "calendar": {k: sorted(v) for k, v in CALENDAR_ALIASES.items()},
        },
        "sample_rows": {
            "transactions": ["L-1001", "TX-1001", "2026-09-08T10:15:00+05:00", "2026-09-08", "TILL-1", "SKU-001", "Knorr Noodles", "Snacks", "2", "180", "360", "130", "sale"],
            "inventory": ["2026-09-08T09:00:00+05:00", "SKU-001", "24", "12", "2026-09-12T09:00:00+05:00", "7", "", "130", "Instant noodles"],
            "footfall": ["EV-1", "2026-09-08T10:10:00+05:00", "in", "ENTRANCE-1", ""],
            "calendar": ["2026-09-08", "true"],
        },
    }


def parse_import_payload(files: list[tuple[str, bytes]], source: str) -> dict[str, ParsedSheet]:
    parsed: dict[str, ParsedSheet] = {}
    for filename, data in files:
        suffix = Path(filename).suffix.lower()
        if suffix == ".xlsx":
            workbook = read_workbook_bytes(data)
            for kind, rows in workbook.items():
                parsed[kind] = parse_sheet(kind, rows, source)
        elif suffix == ".csv":
            kind = normalize_header(Path(filename).stem)
            if kind not in {"transactions", "inventory", "footfall", "calendar"}:
                raise ValueError(f"CSV filename must identify transactions, inventory, footfall, or calendar: {filename}")
            parsed[kind] = parse_sheet(kind, read_csv_bytes(data), source)
        else:
            raise ValueError(f"Unsupported file type: {filename}")
    return parsed


def preview_import(files: list[tuple[str, bytes]], metadata: dict[str, Any]) -> dict[str, Any]:
    source = str(metadata.get("source") or "manual-upload").strip()
    parsed = parse_import_payload(files, source)
    errors = [err for sheet in parsed.values() for err in sheet.errors]
    warnings = [warn for sheet in parsed.values() for warn in sheet.warnings]
    return {
        "ok": not errors,
        "source": source,
        "coverage": {
            "from": metadata.get("coverage_from"),
            "to": metadata.get("coverage_to"),
            "complete": bool(metadata.get("complete", False)),
        },
        "sheets": {
            kind: {
                "rows": len(sheet.rows),
                "headers": sheet.headers,
                "errors": sheet.errors[:50],
                "warnings": sheet.warnings,
                "preview": sheet.rows[:5],
            }
            for kind, sheet in parsed.items()
        },
        "errors": errors[:100],
        "warnings": sorted(set(warnings)),
    }


def commit_import(conn: sqlite3.Connection, files: list[tuple[str, bytes]], metadata: dict[str, Any]) -> dict[str, Any]:
    source = str(metadata.get("source") or "manual-upload").strip()
    coverage_from = metadata.get("coverage_from")
    coverage_to = metadata.get("coverage_to")
    complete = 1 if metadata.get("complete") else 0
    parsed = parse_import_payload(files, source)
    errors = [err for sheet in parsed.values() for err in sheet.errors]
    if errors:
        return {"ok": False, "errors": errors[:100]}
    import_id = str(uuid.uuid4())
    row_count = sum(len(sheet.rows) for sheet in parsed.values())
    now = iso(datetime.now(APP_TZ))
    try:
        with conn:
            kinds = ",".join(sorted(parsed))
            conn.execute(
                "INSERT INTO imports(id, source, kind, coverage_from, coverage_to, complete, row_count, committed_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (import_id, source, kinds, coverage_from, coverage_to, complete, row_count, now),
            )
            if "transactions" in parsed:
                if not all(row.get("line_item_id") for row in parsed["transactions"].rows) and not complete:
                    raise ValueError("Transactions without line_item_id require complete=true for atomic replacement.")
                for row in parsed["transactions"].rows:
                    if not row.get("line_item_id"):
                        conn.execute(
                            "DELETE FROM transactions WHERE source=? AND business_date=? AND till_id=? AND transaction_id=?",
                            (source, row["business_date"], row["till_id"], row["transaction_id"]),
                        )
                    conn.execute(
                        """
                        INSERT INTO transactions(id, import_id, source, line_item_id, transaction_id, business_date, timestamp, till_id, sku_id, sku_name, category, qty, unit_price, total_price, unit_cost, transaction_type, fingerprint)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        ON CONFLICT(fingerprint) DO UPDATE SET import_id=excluded.import_id, timestamp=excluded.timestamp, sku_name=excluded.sku_name, category=excluded.category, qty=excluded.qty, unit_price=excluded.unit_price, total_price=excluded.total_price, unit_cost=excluded.unit_cost, transaction_type=excluded.transaction_type
                        """,
                        (row["id"], import_id, row["source"], row["line_item_id"], row["transaction_id"], row["business_date"], row["timestamp"], row["till_id"], row["sku_id"], row["sku_name"], row["category"], row["qty"], row["unit_price"], row["total_price"], row["unit_cost"], row["transaction_type"], row["fingerprint"]),
                    )
            if "inventory" in parsed:
                for item in parsed["inventory"].rows:
                    snap = item["snapshot"]
                    conn.execute(
                        """
                        INSERT INTO inventory_snapshots(id, import_id, snapshot_at, sku_id, on_hand, unit_cost, comparison_group, lead_time_days, safety_stock_units, fingerprint)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        ON CONFLICT(fingerprint) DO UPDATE SET import_id=excluded.import_id, on_hand=excluded.on_hand, unit_cost=excluded.unit_cost, comparison_group=excluded.comparison_group, lead_time_days=excluded.lead_time_days, safety_stock_units=excluded.safety_stock_units
                        """,
                        (snap["id"], import_id, snap["snapshot_at"], snap["sku_id"], snap["on_hand"], snap["unit_cost"], snap["comparison_group"], snap["lead_time_days"], snap["safety_stock_units"], snap["fingerprint"]),
                    )
                    receipt = item["receipt"]
                    if receipt:
                        conn.execute(
                            "INSERT INTO inventory_receipts(id, import_id, sku_id, quantity, receipt_due_at, fingerprint) VALUES (?, ?, ?, ?, ?, ?) ON CONFLICT(fingerprint) DO UPDATE SET import_id=excluded.import_id, quantity=excluded.quantity, receipt_due_at=excluded.receipt_due_at",
                            (receipt["id"], import_id, receipt["sku_id"], receipt["quantity"], receipt["receipt_due_at"], receipt["fingerprint"]),
                        )
            if "footfall" in parsed:
                for row in parsed["footfall"].rows:
                    conn.execute(
                        "INSERT INTO footfall_events(id, import_id, event_id, timestamp, direction, camera_id, fingerprint) VALUES (?, ?, ?, ?, ?, ?, ?) ON CONFLICT(fingerprint) DO NOTHING",
                        (row["id"], import_id, row["event_id"], row["timestamp"], row["direction"], row["camera_id"], row["fingerprint"]),
                    )
                cameras = sorted({row["camera_id"] for row in parsed["footfall"].rows})
                if coverage_from and coverage_to:
                    for camera in cameras:
                        conn.execute(
                            "INSERT INTO camera_coverage(id, import_id, camera_id, coverage_from, coverage_to) VALUES (?, ?, ?, ?, ?)",
                            (str(uuid.uuid4()), import_id, camera, coverage_from, coverage_to),
                        )
            if "calendar" in parsed:
                for row in parsed["calendar"].rows:
                    conn.execute(
                        "INSERT INTO store_calendar(date, is_open) VALUES (?, ?) ON CONFLICT(date) DO UPDATE SET is_open=excluded.is_open",
                        (row["date"], row["is_open"]),
                    )
    except Exception as exc:
        return {"ok": False, "errors": [{"row": 0, "errors": [str(exc)]}]}
    dataset_version = dataset_version_for(conn)
    job = create_analysis_run(conn, None, None, 7)
    return {"ok": True, "import_id": import_id, "row_count": row_count, "dataset_version": dataset_version, "job_id": job["id"]}


def dataset_version_for(conn: sqlite3.Connection) -> str:
    rows = conn.execute("SELECT id, committed_at, row_count FROM imports ORDER BY committed_at, id").fetchall()
    return stable_hash(*[f"{r['id']}:{r['committed_at']}:{r['row_count']}" for r in rows])[:16]


def date_range_days(date_from: str, date_to: str) -> list[str]:
    start = datetime.fromisoformat(date_from).date()
    end = datetime.fromisoformat(date_to).date()
    return [(start + timedelta(days=i)).isoformat() for i in range((end - start).days + 1)]


def fetch_rows(conn: sqlite3.Connection, sql: str, params: tuple[Any, ...]) -> list[sqlite3.Row]:
    return list(conn.execute(sql, params).fetchall())


def daily_demand(conn: sqlite3.Connection, sku_id: str, end_date: str, max_days: int = 180) -> list[tuple[str, float]]:
    end = datetime.fromisoformat(end_date).date()
    start = end - timedelta(days=max_days - 1)
    rows = fetch_rows(
        conn,
        """
        SELECT business_date, SUM(CASE WHEN transaction_type='sale' THEN qty ELSE 0 END) AS units
        FROM transactions
        WHERE sku_id=? AND business_date BETWEEN ? AND ?
        GROUP BY business_date
        """,
        (sku_id, start.isoformat(), end.isoformat()),
    )
    by_date = {r["business_date"]: float(r["units"] or 0) for r in rows}
    days = [(start + timedelta(days=i)).isoformat() for i in range(max_days)]
    return [(day, max(0.0, by_date.get(day, 0.0))) for day in days]


def complete_calendar(conn: sqlite3.Connection, days: list[str]) -> bool:
    if not days:
        return False
    count = conn.execute(
        f"SELECT COUNT(*) AS c FROM store_calendar WHERE is_open=1 AND date IN ({','.join('?' for _ in days)})",
        tuple(days),
    ).fetchone()["c"]
    return count == len(days)


def wape(actual: list[float], pred: list[float]) -> float | None:
    denom = sum(abs(v) for v in actual)
    if denom == 0:
        return None
    return sum(abs(a - p) for a, p in zip(actual, pred)) / denom


def mae(actual: list[float], pred: list[float]) -> float:
    if not actual:
        return 0.0
    return sum(abs(a - p) for a, p in zip(actual, pred)) / len(actual)


def seasonal_forecast(history: list[float], horizon: int) -> list[float]:
    if len(history) >= 7:
        base = history[-7:]
        return [max(0.0, base[i % 7]) for i in range(horizon)]
    avg = statistics.mean(history[-14:]) if history else 0.0
    return [max(0.0, avg) for _ in range(horizon)]


def moving_average_forecast(history: list[float], horizon: int) -> list[float]:
    if not history:
        return [0.0 for _ in range(horizon)]
    recent = history[-min(28, len(history)) :]
    weekly = seasonal_forecast(history, horizon)
    avg = statistics.mean(recent)
    return [max(0.0, (avg + weekly[i]) / 2) for i in range(horizon)]


def backtest(history: list[float], horizon: int) -> dict[str, Any]:
    eligibility = 56 if horizon == 7 else 120
    cutoffs: list[int] = []
    for cutoff in range(max(28, horizon), len(history) - horizon + 1, horizon):
        if len(cutoffs) >= 6:
            break
        cutoffs.append(cutoff)
    models = {"seasonal_naive": seasonal_forecast, "lag_average": moving_average_forecast}
    results: dict[str, Any] = {}
    for name, fn in models.items():
        actual_all: list[float] = []
        pred_all: list[float] = []
        for cutoff in cutoffs:
            train = history[:cutoff]
            actual = history[cutoff : cutoff + horizon]
            pred = fn(train, horizon)
            actual_all.extend(actual)
            pred_all.extend(pred)
        results[name] = {"wape": wape(actual_all, pred_all), "mae": mae(actual_all, pred_all), "cutoffs": len(cutoffs)}
    usable = len(history) >= eligibility and len(cutoffs) >= 3
    if not usable:
        selected = "recent_weekly_average"
        status = "fallback"
        confidence = "low"
        forecast = moving_average_forecast(history, horizon)
    else:
        selected = sorted(models.keys(), key=lambda key: (float("inf") if results[key]["wape"] is None else results[key]["wape"], 0 if key == "seasonal_naive" else 1))[0]
        score = results[selected]["wape"]
        status = "normal" if score is not None else "fallback"
        if score is None:
            confidence = "low"
        elif score <= 0.20:
            confidence = "high"
        elif score <= 0.40:
            confidence = "medium"
        else:
            confidence = "low"
        forecast = models[selected](history, horizon)
    return {
        "model": selected,
        "status": status,
        "confidence": confidence,
        "history_days": len(history),
        "validation": results,
        "forecast_units": [round(v, 2) for v in forecast],
        "total_units": round(sum(forecast), 2),
    }


def latest_inventory(conn: sqlite3.Connection) -> dict[str, sqlite3.Row]:
    rows = fetch_rows(
        conn,
        """
        SELECT * FROM inventory_snapshots inv
        WHERE snapshot_at = (SELECT MAX(snapshot_at) FROM inventory_snapshots WHERE sku_id=inv.sku_id)
        """,
        (),
    )
    return {r["sku_id"]: r for r in rows}


def transaction_metrics(conn: sqlite3.Connection, date_from: str, date_to: str) -> dict[str, Any]:
    rows = fetch_rows(
        conn,
        "SELECT * FROM transactions WHERE business_date BETWEEN ? AND ?",
        (date_from, date_to),
    )
    gross_sales = sum(r["total_price"] for r in rows if r["transaction_type"] == "sale")
    refunds = abs(sum(r["total_price"] for r in rows if r["transaction_type"] == "refund"))
    sale_lines = [r for r in rows if r["transaction_type"] == "sale"]
    gross_units = sum(r["qty"] for r in sale_lines)
    net_units = sum(r["qty"] for r in rows)
    sale_txs = {(r["source"], r["business_date"], r["till_id"], r["transaction_id"]) for r in sale_lines if r["qty"] > 0}
    baskets: dict[tuple[str, str, str, str], dict[str, float]] = {}
    complete_cost = True
    for r in sale_lines:
        key = (r["source"], r["business_date"], r["till_id"], r["transaction_id"])
        baskets.setdefault(key, {"units": 0.0, "value": 0.0, "cost": 0.0, "cost_lines": 0.0, "lines": 0.0})
        baskets[key]["units"] += r["qty"]
        baskets[key]["value"] += r["total_price"]
        baskets[key]["lines"] += 1
        if r["unit_cost"] is None:
            complete_cost = False
        else:
            baskets[key]["cost"] += r["unit_cost"] * r["qty"]
            baskets[key]["cost_lines"] += 1
    tx_count = len(sale_txs)
    margin_value = None
    margin_status = "unavailable"
    costed_revenue = sum(r["total_price"] for r in sale_lines if r["unit_cost"] is not None)
    if costed_revenue:
        cost_total = sum((r["unit_cost"] or 0) * r["qty"] for r in sale_lines if r["unit_cost"] is not None)
        margin_value = gross_sales - refunds - cost_total
        margin_status = "complete" if complete_cost else "partial"
    return {
        "gross_sales": round(gross_sales, 2),
        "refunds": round(refunds, 2),
        "net_revenue": round(gross_sales - refunds, 2),
        "gross_sold_units": round(gross_units, 2),
        "net_units": round(net_units, 2),
        "distinct_sale_transactions": tx_count,
        "average_basket_value": round((gross_sales / tx_count) if tx_count else 0, 2),
        "items_per_basket": round((gross_units / tx_count) if tx_count else 0, 2),
        "margin": None if margin_value is None else round(margin_value, 2),
        "margin_status": margin_status,
    }


def product_performance(conn: sqlite3.Connection, date_from: str, date_to: str) -> list[dict[str, Any]]:
    rows = fetch_rows(
        conn,
        """
        SELECT sku_id, sku_name, category,
          SUM(CASE WHEN transaction_type='sale' THEN qty ELSE 0 END) AS gross_units,
          SUM(qty) AS net_units,
          SUM(total_price) AS net_revenue,
          SUM(CASE WHEN transaction_type='sale' THEN total_price ELSE 0 END) AS gross_revenue,
          COUNT(DISTINCT source || '|' || business_date || '|' || till_id || '|' || transaction_id) AS txs
        FROM transactions
        WHERE business_date BETWEEN ? AND ?
        GROUP BY sku_id, sku_name, category
        ORDER BY gross_units DESC, net_revenue DESC
        """,
        (date_from, date_to),
    )
    total_units = sum(float(r["gross_units"] or 0) for r in rows)
    return [
        {
            "sku_id": r["sku_id"],
            "sku_name": r["sku_name"],
            "category": r["category"],
            "gross_units": round(float(r["gross_units"] or 0), 2),
            "net_units": round(float(r["net_units"] or 0), 2),
            "net_revenue": round(float(r["net_revenue"] or 0), 2),
            "gross_revenue": round(float(r["gross_revenue"] or 0), 2),
            "transaction_count": int(r["txs"] or 0),
            "sales_velocity_units_per_day": round(float(r["gross_units"] or 0) / max(1, len(date_range_days(date_from, date_to))), 2),
            "unit_share": round((float(r["gross_units"] or 0) / total_units) if total_units else 0, 4),
        }
        for r in rows
    ]


def category_share(products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    totals: dict[str, float] = {}
    for p in products:
        totals[p["category"]] = totals.get(p["category"], 0.0) + p["net_revenue"]
    total = sum(totals.values())
    return [{"category": k, "net_revenue": round(v, 2), "share": round(v / total, 4) if total else 0} for k, v in sorted(totals.items(), key=lambda item: item[1], reverse=True)]


def conversion_metrics(conn: sqlite3.Connection, date_from: str, date_to: str) -> dict[str, Any]:
    entries = conn.execute(
        "SELECT COUNT(*) AS c FROM footfall_events WHERE direction='in' AND date(timestamp) BETWEEN ? AND ?",
        (date_from, date_to),
    ).fetchone()["c"]
    coverage = conn.execute(
        "SELECT COUNT(*) AS c FROM camera_coverage WHERE date(coverage_from) <= ? AND date(coverage_to) >= ?",
        (date_to, date_from),
    ).fetchone()["c"]
    txs = transaction_metrics(conn, date_from, date_to)["distinct_sale_transactions"]
    ratio = None if not entries else txs / entries
    flags = []
    if not coverage:
        flags.append("missing_camera_coverage")
    if entries == 0:
        flags.append("zero_entries")
    if ratio is not None and ratio > 1:
        flags.append("conversion_above_100_percent")
    daily = []
    for day in date_range_days(date_from, date_to):
        day_entries = conn.execute("SELECT COUNT(*) AS c FROM footfall_events WHERE direction='in' AND date(timestamp)=?", (day,)).fetchone()["c"]
        day_txs = transaction_metrics(conn, day, day)["distinct_sale_transactions"]
        daily.append({"period": day, "entries": day_entries, "transactions": day_txs, "conversion": None if not day_entries else round(day_txs / day_entries, 4)})
    return {"entries": entries, "transactions": txs, "conversion": None if ratio is None else round(ratio, 4), "flags": flags, "buckets": daily, "definition": "completed sale transactions divided by entry events; not identified visitor purchase rate"}


def trend_for_product(conn: sqlite3.Connection, sku_id: str, date_from: str, date_to: str) -> dict[str, Any]:
    days = len(date_range_days(date_from, date_to))
    start = datetime.fromisoformat(date_from).date()
    prior_to = (start - timedelta(days=1)).isoformat()
    prior_from = (start - timedelta(days=days)).isoformat()
    current = conn.execute("SELECT SUM(CASE WHEN transaction_type='sale' THEN qty ELSE 0 END) AS u FROM transactions WHERE sku_id=? AND business_date BETWEEN ? AND ?", (sku_id, date_from, date_to)).fetchone()["u"] or 0
    prior = conn.execute("SELECT SUM(CASE WHEN transaction_type='sale' THEN qty ELSE 0 END) AS u FROM transactions WHERE sku_id=? AND business_date BETWEEN ? AND ?", (sku_id, prior_from, prior_to)).fetchone()["u"] or 0
    change = None if prior == 0 else (current - prior) / prior
    return {"current_units": round(current, 2), "prior_units": round(prior, 2), "change_pct": None if change is None else round(change, 4), "prior_period": {"from": prior_from, "to": prior_to}}


def recommendation_for(conn: sqlite3.Connection, product: dict[str, Any], forecast: dict[str, Any], horizon: int, analysis_to: str) -> dict[str, Any]:
    inv = latest_inventory(conn).get(product["sku_id"])
    warnings = []
    if inv is None:
        return {"sku_id": product["sku_id"], "sku_name": product["sku_name"], "status": "MONITOR", "confidence": "low", "rule": "No inventory snapshot is available.", "warnings": ["missing_inventory"], "suggested_order_qty": 0}
    snapshot_at = parse_timestamp(inv["snapshot_at"])
    age_hours = (datetime.now(APP_TZ) - snapshot_at).total_seconds() / 3600
    stale = age_hours > 24
    if stale:
        warnings.append("stale_inventory_snapshot")
    if inv["on_hand"] <= 0:
        warnings.append("snapshot_records_zero_stock")
    lead = int(inv["lead_time_days"] or 7)
    history = [v for _, v in daily_demand(conn, product["sku_id"], analysis_to, 180)]
    safety = inv["safety_stock_units"]
    if safety is None:
        std = statistics.pstdev(history[-28:]) if len(history) >= 2 else 0.0
        safety = math.ceil(1.65 * std * math.sqrt(lead))
    receipts_rows = fetch_rows(conn, "SELECT quantity, receipt_due_at FROM inventory_receipts WHERE sku_id=?", (product["sku_id"],))
    undated = sum(float(r["quantity"]) for r in receipts_rows if r["receipt_due_at"] is None)
    if undated:
        warnings.append("undated_on_order_cannot_prevent_stockout")
    dated_receipts: dict[int, float] = {}
    snapshot_date = snapshot_at.date()
    for r in receipts_rows:
        if r["receipt_due_at"]:
            offset = (parse_timestamp(r["receipt_due_at"]).date() - snapshot_date).days
            if offset >= 0:
                dated_receipts[offset] = dated_receipts.get(offset, 0.0) + float(r["quantity"])
    days_needed = max(lead + horizon, 30)
    daily_fc = forecast["forecast_units"]
    while len(daily_fc) < days_needed:
        daily_fc.extend(daily_fc[-7:] if len(daily_fc) >= 7 else [0.0])
    balance = float(inv["on_hand"])
    projected = []
    stockout_day = None
    below_safety_before_arrival = False
    for day_idx in range(days_needed):
        balance += dated_receipts.get(day_idx, 0.0)
        balance -= daily_fc[day_idx]
        projected.append({"day": day_idx + 1, "balance": round(balance, 2), "receipt": round(dated_receipts.get(day_idx, 0.0), 2), "forecast_units": round(daily_fc[day_idx], 2)})
        if balance < 0 and stockout_day is None:
            stockout_day = day_idx + 1
        if day_idx < lead and balance < float(safety):
            below_safety_before_arrival = True
    receipts_due_period = sum(qty for day, qty in dated_receipts.items() if day < lead + horizon)
    demand_period = sum(daily_fc[: lead + horizon])
    suggested = math.ceil(max(0.0, demand_period + float(safety) - float(inv["on_hand"]) - receipts_due_period))
    trend = trend_for_product(conn, product["sku_id"], (datetime.fromisoformat(analysis_to).date() - timedelta(days=horizon - 1)).isoformat(), analysis_to)
    coverage_days = (float(inv["on_hand"]) + receipts_due_period) / max(0.01, sum(daily_fc[:30]) / 30)
    declining = trend["change_pct"] is not None and trend["change_pct"] <= -0.20
    if stale or forecast["confidence"] == "low":
        status = "MONITOR"
        rule = "Inventory is stale or forecast confidence is low; no actionable order quantity is issued."
    elif below_safety_before_arrival:
        status = "RESTOCK"
        rule = "Projected stock before a new order can arrive falls below safety stock."
    elif coverage_days > 30 and declining:
        status = "HOLD"
        rule = "Stock covers more than 30 forecast days and comparable demand declined by at least 20%."
        suggested = 0
    else:
        status = "MONITOR"
        rule = "No restock or hold rule boundary was crossed."
    if stockout_day is not None and stockout_day < lead:
        warnings.append("predicted_stockout_before_replenishment_arrival")
    return {
        "sku_id": product["sku_id"],
        "sku_name": product["sku_name"],
        "status": status,
        "confidence": forecast["confidence"],
        "rule": rule,
        "suggested_order_qty": suggested if not stale and forecast["confidence"] != "low" else 0,
        "snapshot_at": inv["snapshot_at"],
        "on_hand": inv["on_hand"],
        "lead_time_days": lead,
        "safety_stock_units": safety,
        "coverage_days": round(coverage_days, 1),
        "trend": trend,
        "projected_balances": projected[: lead + horizon],
        "incoming_receipts": [{"quantity": r["quantity"], "receipt_due_at": r["receipt_due_at"]} for r in receipts_rows],
        "warnings": warnings,
    }


def anomalies(products: list[dict[str, Any]], conversion: dict[str, Any]) -> list[dict[str, Any]]:
    result = []
    for p in products:
        if p["gross_units"] == 0:
            result.append({"scope": "sku", "scope_id": p["sku_id"], "severity": "info", "message": f"{p['sku_name']} has zero gross sold units in this period."})
    if "conversion_above_100_percent" in conversion["flags"]:
        result.append({"scope": "store", "scope_id": "conversion", "severity": "warning", "message": "Purchase conversion is above 100%; this usually means multiple transactions per visitor, mismatched timing, or missing entry events."})
    return result[:20]


def create_evidence(conn: sqlite3.Connection, run_id: str, result: dict[str, Any]) -> None:
    conn.execute("DELETE FROM evidence_fts WHERE id IN (SELECT id FROM evidence WHERE run_id=?)", (run_id,))
    conn.execute("DELETE FROM evidence WHERE run_id=?", (run_id,))
    items: list[tuple[str, str, str, str, str, str, str]] = []
    kpi = result["kpis"]
    text = f"Store KPIs from {result['period']['from']} to {result['period']['to']}: net revenue PKR {kpi['net_revenue']}, gross sold units {kpi['gross_sold_units']}, sale transactions {kpi['distinct_sale_transactions']}, average basket PKR {kpi['average_basket_value']}."
    items.append((f"ev-{run_id[:8]}-store", "store", "all", result["period"]["from"], result["period"]["to"], text, json.dumps(kpi)))
    conv = result["conversion"]
    items.append((f"ev-{run_id[:8]}-conversion", "conversion", "all", result["period"]["from"], result["period"]["to"], f"Purchase conversion from {result['period']['from']} to {result['period']['to']} was {conv['conversion']} using {conv['transactions']} sale transactions and {conv['entries']} entry events. Flags: {', '.join(conv['flags']) or 'none'}.", json.dumps(conv)))
    for p in result["products"][:50]:
        rec = next((r for r in result["recommendations"] if r["sku_id"] == p["sku_id"]), None)
        fc = result["forecasts"].get(p["sku_id"], {})
        text = f"SKU {p['sku_id']} {p['sku_name']} sold {p['gross_units']} gross units and PKR {p['net_revenue']} net revenue from {result['period']['from']} to {result['period']['to']}. Forecast {fc.get('total_units')} units over {result['horizon']} days with {fc.get('confidence')} confidence. Recommendation {rec.get('status') if rec else 'MONITOR'}: {rec.get('rule') if rec else 'no inventory evidence'}."
        items.append((f"ev-{run_id[:8]}-{stable_hash(p['sku_id'])[:8]}", "sku", p["sku_id"], result["period"]["from"], result["period"]["to"], text, json.dumps({"product": p, "forecast": fc, "recommendation": rec})))
    for item in items:
        conn.execute("INSERT INTO evidence(id, run_id, scope_type, scope_id, period_from, period_to, text, payload_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?)", (item[0], run_id, *item[1:]))
        conn.execute("INSERT INTO evidence_fts(id, text) VALUES (?, ?)", (item[0], item[5]))


def run_analysis(conn: sqlite3.Connection, date_from: str, date_to: str, horizon: int) -> dict[str, Any]:
    products = product_performance(conn, date_from, date_to)
    kpis = transaction_metrics(conn, date_from, date_to)
    conversion = conversion_metrics(conn, date_from, date_to)
    days = date_range_days(date_from, date_to)
    calendar_complete = complete_calendar(conn, days)
    forecasts: dict[str, Any] = {}
    recs: list[dict[str, Any]] = []
    for product in products:
        hist = [v for _, v in daily_demand(conn, product["sku_id"], date_to, 180)]
        fc = backtest(hist, horizon)
        if not calendar_complete:
            fc["confidence"] = "low"
            fc["status"] = "coverage_unknown"
        forecasts[product["sku_id"]] = fc
        recs.append(recommendation_for(conn, product, fc, horizon, date_to))
    quality_flags = []
    if not calendar_complete:
        quality_flags.append("missing_or_incomplete_opening_calendar")
    quality_flags.extend(conversion["flags"])
    return {
        "run_id": None,
        "dataset_version": dataset_version_for(conn),
        "period": {"from": date_from, "to": date_to},
        "horizon": horizon,
        "currency": "PKR",
        "timezone": str(APP_TZ),
        "kpis": kpis,
        "products": products,
        "categories": category_share(products),
        "conversion": conversion,
        "forecasts": forecasts,
        "recommendations": recs,
        "anomalies": anomalies(products, conversion),
        "quality_flags": sorted(set(quality_flags)),
        "generated_at": iso(datetime.now(APP_TZ)),
    }


def create_analysis_run(conn: sqlite3.Connection, date_from: str | None, date_to: str | None, horizon: int, status: str = "completed") -> dict[str, Any]:
    if date_from is None or date_to is None:
        latest = conn.execute("SELECT MAX(business_date) AS d FROM transactions").fetchone()["d"]
        if latest:
            date_to = latest
            date_from = (datetime.fromisoformat(latest).date() - timedelta(days=6)).isoformat()
        else:
            today = datetime.now(APP_TZ).date().isoformat()
            date_from = date_to = today
    run_id = str(uuid.uuid4())
    conn.execute(
        "INSERT INTO analysis_runs(id, dataset_version, date_from, date_to, horizon, status, created_at, settings_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (run_id, dataset_version_for(conn), date_from, date_to, horizon, "running" if status == "completed" else status, iso(datetime.now(APP_TZ)), "{}"),
    )
    conn.commit()
    if status == "completed":
        result = run_analysis(conn, date_from, date_to, horizon)
        result["run_id"] = run_id
        with conn:
            conn.execute("UPDATE analysis_runs SET status='completed', completed_at=?, result_json=? WHERE id=?", (iso(datetime.now(APP_TZ)), json.dumps(result), run_id))
            create_evidence(conn, run_id, result)
    return get_job(conn, run_id)


def get_job(conn: sqlite3.Connection, job_id: str) -> dict[str, Any]:
    row = conn.execute("SELECT * FROM analysis_runs WHERE id=?", (job_id,)).fetchone()
    if not row:
        raise KeyError("job not found")
    return dict(row)


def latest_analytics(conn: sqlite3.Connection, date_from: str, date_to: str, horizon: int) -> dict[str, Any]:
    row = conn.execute(
        "SELECT * FROM analysis_runs WHERE date_from=? AND date_to=? AND horizon=? AND status='completed' ORDER BY completed_at DESC LIMIT 1",
        (date_from, date_to, horizon),
    ).fetchone()
    if row and row["result_json"]:
        return json.loads(row["result_json"])
    job = create_analysis_run(conn, date_from, date_to, horizon)
    return json.loads(job["result_json"]) if job.get("result_json") else run_analysis(conn, date_from, date_to, horizon)


def default_window(conn: sqlite3.Connection) -> tuple[str, str]:
    latest = conn.execute("SELECT MAX(business_date) AS d FROM transactions").fetchone()["d"]
    if not latest:
        today = datetime.now(APP_TZ).date()
        return today.isoformat(), today.isoformat()
    end = datetime.fromisoformat(latest).date()
    return (end - timedelta(days=6)).isoformat(), end.isoformat()


def search_evidence(conn: sqlite3.Connection, question: str, date_from: str | None = None, date_to: str | None = None, sku_ids: list[str] | None = None, run_id: str | None = None) -> list[dict[str, Any]]:
    clauses = []
    params: list[Any] = []
    if run_id:
        clauses.append("e.run_id=?")
        params.append(run_id)
    if date_from:
        clauses.append("e.period_to>=?")
        params.append(date_from)
    if date_to:
        clauses.append("e.period_from<=?")
        params.append(date_to)
    if sku_ids:
        clauses.append(f"(e.scope_type!='sku' OR e.scope_id IN ({','.join('?' for _ in sku_ids)}))")
        params.extend(sku_ids)
    where = (" AND " + " AND ".join(clauses)) if clauses else ""
    tokens = [t for t in re.findall(r"[A-Za-z0-9]+", question) if len(t) > 2]
    query = " OR ".join(tokens[:8]) or question
    try:
        rows = fetch_rows(
            conn,
            f"SELECT e.* FROM evidence_fts f JOIN evidence e ON e.id=f.id WHERE evidence_fts MATCH ? {where} LIMIT 8",
            (query, *params),
        )
    except sqlite3.OperationalError:
        rows = fetch_rows(conn, f"SELECT e.* FROM evidence e WHERE 1=1 {where} LIMIT 8", tuple(params))
    return [{"id": r["id"], "scope_type": r["scope_type"], "scope_id": r["scope_id"], "period": {"from": r["period_from"], "to": r["period_to"]}, "text": r["text"], "payload": json.loads(r["payload_json"])} for r in rows]


def answer_query(conn: sqlite3.Connection, question: str, filters: dict[str, Any]) -> dict[str, Any]:
    date_from = filters.get("from")
    date_to = filters.get("to")
    if not date_from or not date_to:
        date_from, date_to = default_window(conn)
    horizon = int(filters.get("horizon") or 7)
    analytics = latest_analytics(conn, date_from, date_to, horizon)
    q = question.lower()
    if any(word in q for word in ["top", "best", "ranking", "rank"]):
        evidence = []
        lines = []
        for p in analytics["products"][:5]:
            ev = search_evidence(conn, p["sku_id"], date_from, date_to, [p["sku_id"]], analytics["run_id"])[:1]
            evidence.extend(ev)
            lines.append(f"{p['sku_name']} sold {p['gross_units']} units and generated PKR {p['net_revenue']}.")
        return {"answer": " ".join(lines) if lines else "There is not enough product evidence for that ranking.", "mode": "structured", "citations": evidence}
    if "restock" in q or "hold" in q or "recommend" in q:
        recs = [r for r in analytics["recommendations"] if r["status"] in {"RESTOCK", "HOLD"}] or analytics["recommendations"][:5]
        evidence = []
        lines = []
        for r in recs[:5]:
            ev = search_evidence(conn, r["sku_id"], date_from, date_to, [r["sku_id"]], analytics["run_id"])[:1]
            evidence.extend(ev)
            lines.append(f"{r['sku_name']}: {r['status']} because {r['rule']} Suggested order quantity: {r.get('suggested_order_qty', 0)}.")
        return {"answer": " ".join(lines) if lines else "There is not enough inventory evidence for a recommendation.", "mode": "structured", "citations": evidence}
    evidence = search_evidence(conn, question, date_from, date_to, filters.get("sku_ids"), analytics["run_id"])
    if not evidence:
        return {"answer": "I do not have enough internal evidence for that question in the selected period.", "mode": "insufficient_evidence", "citations": []}
    return {"answer": " ".join(item["text"] for item in evidence[:3]), "mode": "extractive", "citations": evidence[:5]}
